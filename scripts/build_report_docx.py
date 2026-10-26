"""Build the formatted project report DOCX from report/final_report.md."""

from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "report" / "final_report.md"
OUTPUT = ROOT / "report" / "AI_HVAC_Fault_Detection_Report.docx"

BODY_FONT = "Liberation Serif"
HEADING_FONT = "Liberation Sans"
MONO_FONT = "Liberation Mono"
NAVY = "17365D"
PALE_BLUE = "EDF3F8"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "666666"


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), LIGHT_GRAY)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), BODY_FONT)
    r_pr.extend([r_fonts, color, underline])
    new_run.append(r_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(
    r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))"
)


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, BODY_FONT)
        token = match.group(0)
        if token.startswith("["):
            link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_run_font(run, BODY_FONT)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, MONO_FONT, 9.5)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            set_run_font(run, BODY_FONT)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, BODY_FONT)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, HEADING_FONT, 9)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)


def set_repeat_and_keep(paragraph, keep_next: bool = False) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    widow = OxmlElement("w:widowControl")
    widow.set(qn("w:val"), "1")
    p_pr.append(widow)
    if keep_next:
        keep = OxmlElement("w:keepNext")
        keep.set(qn("w:val"), "1")
        p_pr.append(keep)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = styles["Title"]
    title.font.name = HEADING_FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), HEADING_FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), HEADING_FONT)
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 17, 16, 8),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = styles[name]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = BODY_FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)


def add_cover(document: Document) -> None:
    for _ in range(4):
        document.add_paragraph()
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Artificial Intelligence for HVAC Fault Detection and Diagnosis")
    set_run_font(run, HEADING_FONT, 26)
    run.bold = True

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("Industry Research Project Report")
    set_run_font(run, HEADING_FONT, 16)

    details = [
        "Smart Facilities SA",
        "Group C262F-B106-D01",
        "27 October 2026",
    ]
    for value in details:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(value)
        set_run_font(run, BODY_FONT, 12)

    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.space_before = Pt(34)
    run = status.add_run("Draft for final team verification")
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    set_run_font(run, BODY_FONT, 10.5)
    document.add_page_break()


def add_contents(document: Document, lines: list[str]) -> None:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(12)
    heading.paragraph_format.keep_with_next = False
    run = heading.add_run("Contents")
    set_run_font(run, HEADING_FONT, 18)
    run.bold = True
    for line in lines:
        match = re.match(r"^(#{1,2})\s+(.+)$", line)
        if not match:
            continue
        label = match.group(2).strip()
        if label == "Artificial Intelligence for HVAC Fault Detection and Diagnosis":
            continue
        if label == "Submission verification checklist":
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0 if len(match.group(1)) == 1 else 0.25)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(label)
        set_run_font(run, HEADING_FONT if len(match.group(1)) == 1 else BODY_FONT, 10)
        run.bold = len(match.group(1)) == 1 or label == "Executive Summary"
    document.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    table_lines = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        table_lines.append(lines[index].strip())
        index += 1
    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            text = row[column_index] if column_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            if re.fullmatch(r"[\d.,%—-]+|\*\*\[[^]]+\]\*\*", text):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, text)
            for run in paragraph.runs:
                set_run_font(run, HEADING_FONT if row_index == 0 else BODY_FONT, 9 if column_count >= 6 else 9.5)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(document: Document, alt: str, relative_path: str) -> None:
    image_path = (SOURCE.parent / relative_path).resolve()
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.25))
    inline = document.inline_shapes[-1]._inline
    doc_pr = inline.docPr
    doc_pr.set("descr", alt)
    doc_pr.set("title", alt[:80])


def add_body(document: Document, lines: list[str]) -> None:
    started = False
    index = 0
    first_body_heading = True
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped == "## Executive Summary":
            started = True
            paragraph = document.add_paragraph("Executive Summary", style="Heading 1")
            set_repeat_and_keep(paragraph, True)
            first_body_heading = False
            index += 1
            continue
        if not started:
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        image = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", stripped)
        if image:
            add_image(document, image.group(1), image.group(2))
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            label = heading.group(2).strip()
            paragraph = document.add_paragraph(label, style=f"Heading {min(level, 3)}")
            set_repeat_and_keep(paragraph, True)
            first_body_heading = False
            index += 1
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.right_indent = Inches(0.25)
            paragraph.paragraph_format.space_after = Pt(10)
            add_inline(paragraph, stripped.lstrip("> "))
            for run in paragraph.runs:
                run.italic = True
            set_repeat_and_keep(paragraph)
            index += 1
            continue
        if re.match(r"^-\s+", stripped):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(6)
            add_inline(paragraph, re.sub(r"^-\s+", "", stripped))
            set_repeat_and_keep(paragraph)
            index += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(6)
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            set_repeat_and_keep(paragraph)
            index += 1
            continue

        paragraph = document.add_paragraph()
        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            paragraph.style = document.styles["Caption"]
            add_inline(paragraph, stripped[1:-1])
            for run in paragraph.runs:
                run.italic = True
        elif stripped.startswith("`") and stripped.endswith("`"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(stripped[1:-1])
            set_run_font(run, MONO_FONT, 9.5)
        else:
            add_inline(paragraph, stripped)
        set_repeat_and_keep(paragraph)
        index += 1


def configure_page(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.95)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("AI for HVAC Fault Detection and Diagnosis")
        set_run_font(run, HEADING_FONT, 8.5)
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)
        add_page_number(section.footer.paragraphs[0])
        section.first_page_header.paragraphs[0].text = ""
        section.first_page_footer.paragraphs[0].text = ""


def set_document_metadata(document: Document) -> None:
    properties = document.core_properties
    properties.title = "Artificial Intelligence for HVAC Fault Detection and Diagnosis"
    properties.subject = "Industry Research Project report for Smart Facilities SA"
    properties.author = "Group C262F-B106-D01"
    properties.keywords = "HVAC, fault detection, anomaly detection, XGBoost, facilities management"


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_styles(document)
    configure_page(document)
    set_document_metadata(document)
    add_cover(document)
    add_contents(document, lines)
    add_body(document, lines)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
